from pathlib import Path

import pytest
from PIL import Image

from learning_ext import document_preview as preview


@pytest.fixture(autouse=True)
def cache(monkeypatch, tmp_path):
    monkeypatch.setattr(preview, "CACHE", tmp_path / "cache")


def test_images_and_multiframe_tiff_preview_without_external_apps(tmp_path):
    for suffix in ("png", "jpg", "webp", "bmp", "gif", "tiff"):
        path = tmp_path / ("example." + suffix)
        Image.new("RGB", (240, 320), "skyblue").save(path)
        doc = preview.prepare_file(path, path.name)
        html, number, count = preview.render_page(doc)
        assert "data:image/png;base64," in html
        assert number == 1 and count == "1 / 1"
    path = tmp_path / "pages.tiff"
    Image.new("RGB", (30, 40), "red").save(path, save_all=True, append_images=[Image.new("RGB", (30, 40), "blue")])
    doc = preview.prepare_file(path, path.name)
    assert doc["pages"] == 2
    assert preview.render_page(doc, 99)[1] == 2


def test_unsupported_and_missing_files_are_not_executed(tmp_path):
    path = tmp_path / "code.html"
    path.write_text("<script>alert(1)</script>")
    assert preview.prepare_file(path, path.name) is None
    with pytest.raises(ValueError):
        preview.render_page({"path": str(path), "pages": 1})


def test_save_originals_is_private_idempotent_and_rejects_outside_uploads(tmp_path, monkeypatch):
    import uuid
    from types import SimpleNamespace
    from sqlalchemy import Column, String, Integer, JSON, create_engine
    from sqlalchemy.orm import declarative_base
    import ktem.db.engine as database

    base = declarative_base()
    class Source(base):
        __tablename__ = "preview_sources"
        id = Column(String, primary_key=True, default=lambda: str(uuid.uuid4()))
        name = Column(String)
        path = Column(String)
        size = Column(Integer)
        user = Column(String)
        note = Column(JSON)

    engine = create_engine("sqlite:///:memory:")
    base.metadata.create_all(engine)
    monkeypatch.setattr(database, "engine", engine)
    upload = tmp_path / "uploads"
    upload.mkdir()
    monkeypatch.setenv("GRADIO_TEMP_DIR", str(upload))
    index = SimpleNamespace(_resources={"Source": Source}, _fs_path=tmp_path / "stored", config={"private": True})
    path = upload / "picture.png"
    Image.new("RGB", (20, 20), "blue").save(path)
    first = preview.save_originals(index, [path], "alice")
    assert preview.save_originals(index, [path], "alice") == first
    second = preview.save_originals(index, [path], "bob")
    assert first[0] != second[0]
    assert preview.source_file(index, first[0], "alice")[1] == path.name
    with pytest.raises(ValueError, match="访问权限"):
        preview.source_file(index, first[0], "bob")
    Image.new("RGB", (20, 20), "red").save(path)
    with pytest.raises(ValueError, match="同名资料"):
        preview.save_originals(index, [path], "alice")
    outside = tmp_path / "outside.png"
    Image.new("RGB", (20, 20)).save(outside)
    with pytest.raises(ValueError, match="上传窗口"):
        preview.save_originals(index, [outside], "alice")


@pytest.mark.skipif(not all(p.exists() for p in preview.engine_paths()), reason="Bundled Windows renderers unavailable")
def test_bundled_libreoffice_word_and_sumatra_pdf_render_real_pages(tmp_path):
    from docx import Document
    document = Document()
    document.add_heading("Word preview acceptance", 0)
    document.add_paragraph("The original layout stays inside Learn Everything.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Platform"
    table.cell(0, 1).text = "Model"
    table.cell(1, 0).text = "Saved connection"
    table.cell(1, 1).text = "Switch in chat"
    document.add_page_break()
    document.add_heading("Second page", 1)
    path = tmp_path / "word-preview.docx"
    document.save(path)
    doc = preview.prepare_file(path, path.name)
    assert doc["pages"] == 2
    first = preview.render_page(doc, 1)
    second = preview.render_page(doc, 2)
    assert first[0] != second[0]
    assert "LibreOffice" in first[0]
    pdf = Path(doc["path"])
    standalone = preview.prepare_file(pdf, "standalone.pdf")
    assert preview.render_page(standalone, 2)[2] == "2 / 2"
    assert preview.prepare_file(path, path.name)["path"] == doc["path"]
